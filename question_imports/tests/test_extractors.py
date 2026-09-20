from io import BytesIO

import pymupdf
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from question_imports.extractors import ExtractionError, extract_text


def upload(name, data):
    return SimpleUploadedFile(name, data)


def test_extracts_docx_paragraphs_in_order():
    document = Document()
    document.add_paragraph("QUESTION: Docx?")
    document.add_paragraph("OPTION: Yes")
    stream = BytesIO()
    document.save(stream)
    assert extract_text(upload("questions.docx", stream.getvalue())).splitlines()[:2] == [
        "QUESTION: Docx?",
        "OPTION: Yes",
    ]


def test_extracts_pdf_pages_in_order():
    document = pymupdf.open()
    for text in ("QUESTION: First?", "QUESTION: Second?"):
        page = document.new_page()
        page.insert_text((72, 72), text)
    value = extract_text(upload("questions.pdf", document.tobytes()))
    assert value.index("First") < value.index("Second")


def test_rejects_empty_pdf_corrupt_docx_and_unknown_extension():
    empty_pdf = pymupdf.open()
    empty_pdf.new_page()
    with pytest.raises(ExtractionError, match="OCR"):
        extract_text(upload("blank.pdf", empty_pdf.tobytes()))
    with pytest.raises(ExtractionError, match="Could not read"):
        extract_text(upload("broken.docx", b"not a zip"))
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_text(upload("questions.csv", b"QUESTION: No"))


def test_rejects_invalid_utf8_and_oversized_upload():
    with pytest.raises(ExtractionError, match="Could not read"):
        extract_text(upload("bad.txt", b"\xff"))
    with pytest.raises(ExtractionError, match="5 MB"):
        extract_text(upload("large.md", b"x" * (5 * 1024 * 1024 + 1)))
